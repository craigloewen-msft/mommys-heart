import os
import hashlib
from pathlib import Path

from docx import Document
import chromadb
from openai import AzureOpenAI

from app.config import settings

client = AzureOpenAI(
    azure_endpoint=settings.AZURE_OPENAI_ENDPOINT,
    api_key=settings.AZURE_OPENAI_API_KEY,
    api_version=settings.AZURE_OPENAI_API_VERSION,
)

CHUNK_SIZE = 500  # approximate tokens (chars / 4)
CHUNK_OVERLAP = 100


def _get_chroma_collection() -> chromadb.Collection:
    chroma_client = chromadb.PersistentClient(path=settings.CHROMA_DIR)
    return chroma_client.get_or_create_collection(
        name="documents",
        metadata={"hnsw:space": "cosine"},
    )


def _extract_text_from_docx(filepath: str) -> list[dict]:
    """Extract paragraphs and table content from a .docx file with heading context."""
    doc = Document(filepath)
    sections: list[dict] = []
    current_heading = "Introduction"
    filename = os.path.basename(filepath)

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        if para.style and para.style.name and para.style.name.startswith("Heading"):
            current_heading = text
            continue

        sections.append({
            "text": text,
            "heading": current_heading,
            "filename": filename,
        })

    # Extract text from tables
    for table in doc.tables:
        rows_text = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                rows_text.append(" | ".join(cells))
        if rows_text:
            sections.append({
                "text": "\n".join(rows_text),
                "heading": current_heading,
                "filename": filename,
            })

    return sections


def _chunk_sections(sections: list[dict]) -> list[dict]:
    """Group paragraphs into chunks of approximately CHUNK_SIZE tokens."""
    chunks: list[dict] = []
    current_chunk: list[dict] = []  # list of {text, heading} dicts
    current_chunk_len = 0
    current_filename = sections[0]["filename"] if sections else ""

    for section in sections:
        text = section["text"]
        text_tokens = len(text) // 4  # rough token estimate

        if current_chunk_len + text_tokens > CHUNK_SIZE and current_chunk:
            headings = sorted({s["heading"] for s in current_chunk})
            chunks.append({
                "text": "\n\n".join(s["text"] for s in current_chunk),
                "heading": ", ".join(headings),
                "filename": current_filename,
            })
            # Overlap: keep the last paragraph(s) up to CHUNK_OVERLAP tokens
            overlap: list[dict] = []
            overlap_len = 0
            for s in reversed(current_chunk):
                s_len = len(s["text"]) // 4
                if overlap_len + s_len > CHUNK_OVERLAP:
                    break
                overlap.insert(0, s)
                overlap_len += s_len
            current_chunk = overlap
            current_chunk_len = overlap_len

        current_chunk.append({"text": text, "heading": section["heading"]})
        current_chunk_len += text_tokens
        current_filename = section["filename"]

    if current_chunk:
        headings = sorted({s["heading"] for s in current_chunk})
        chunks.append({
            "text": "\n\n".join(s["text"] for s in current_chunk),
            "heading": ", ".join(headings),
            "filename": current_filename,
        })

    return chunks


def _embed_texts(texts: list[str]) -> list[list[float]]:
    """Get embeddings from Azure OpenAI in batches."""
    all_embeddings: list[list[float]] = []
    batch_size = 16
    for i in range(0, len(texts), batch_size):
        batch = texts[i : i + batch_size]
        response = client.embeddings.create(
            input=batch,
            model=settings.AZURE_OPENAI_EMBEDDING_DEPLOYMENT,
        )
        all_embeddings.extend([item.embedding for item in response.data])
    return all_embeddings


def _compute_docs_hash(docs_dir: str) -> str:
    """Compute a hash of all .docx files to detect changes."""
    hasher = hashlib.md5()
    for filepath in sorted(Path(docs_dir).glob("*.docx")):
        hasher.update(filepath.name.encode())
        hasher.update(str(filepath.stat().st_mtime_ns).encode())
    return hasher.hexdigest()


def ingest_documents(force: bool = False) -> dict:
    """Parse, chunk, embed, and store all .docx files. Returns ingestion stats."""
    docs_dir = settings.DOCS_DIR
    collection = _get_chroma_collection()

    # Check if we need to re-ingest
    current_hash = _compute_docs_hash(docs_dir)
    existing_meta = collection.metadata or {}
    if not force and existing_meta.get("docs_hash") == current_hash and collection.count() > 0:
        return {"status": "skipped", "reason": "documents unchanged", "chunks": collection.count()}

    all_chunks: list[dict] = []
    docx_files = list(Path(docs_dir).glob("*.docx"))

    for filepath in docx_files:
        sections = _extract_text_from_docx(str(filepath))
        if sections:
            chunks = _chunk_sections(sections)
            all_chunks.extend(chunks)

    if not all_chunks:
        return {"status": "error", "reason": "no content found in documents"}

    # Embed all chunks BEFORE deleting the old collection (safe replacement)
    texts = [c["text"] for c in all_chunks]
    embeddings = _embed_texts(texts)

    ids = [f"chunk_{i}" for i in range(len(all_chunks))]
    metadatas = [{"filename": c["filename"], "heading": c["heading"]} for c in all_chunks]

    # Only delete+recreate after embeddings succeed
    chroma_client = chromadb.PersistentClient(path=settings.CHROMA_DIR)
    chroma_client.delete_collection("documents")
    collection = chroma_client.create_collection(
        name="documents",
        metadata={"hnsw:space": "cosine", "docs_hash": current_hash},
    )

    collection.add(
        ids=ids,
        embeddings=embeddings,
        documents=texts,
        metadatas=metadatas,
    )

    return {
        "status": "success",
        "files_processed": len(docx_files),
        "chunks_created": len(all_chunks),
    }
