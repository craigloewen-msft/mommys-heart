import os
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from app.config import settings


def slugify(text: str) -> str:
    """Convert heading text to a URL-safe anchor slug."""
    slug = text.lower().strip()
    slug = re.sub(r'[^\w\s-]', '', slug)
    slug = re.sub(r'[\s_]+', '-', slug)
    slug = re.sub(r'-+', '-', slug).strip('-')
    return slug


def _para_to_html(para) -> str:
    """Convert a paragraph to HTML, preserving bold/italic/underline runs."""
    parts = []
    for run in para.runs:
        text = run.text
        if not text:
            continue
        # Escape HTML entities
        text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        if run.bold:
            text = f'<strong>{text}</strong>'
        if run.italic:
            text = f'<em>{text}</em>'
        if run.underline:
            text = f'<u>{text}</u>'
        parts.append(text)
    return ''.join(parts)


def _table_to_html(table) -> str:
    """Convert a docx table to an HTML table."""
    rows_html = []
    for i, row in enumerate(table.rows):
        cells_html = []
        for cell in row.cells:
            cell_text = cell.text.strip().replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            tag = 'th' if i == 0 else 'td'
            cells_html.append(f'<{tag}>{cell_text}</{tag}>')
        rows_html.append(f'<tr>{"".join(cells_html)}</tr>')
    return f'<table>{"".join(rows_html)}</table>'


def render_docx_to_html(filename: str) -> str | None:
    """Render a .docx file as a styled HTML page with heading anchors."""
    filepath = Path(settings.DOCS_DIR) / filename
    if not filepath.exists() or not filepath.suffix == '.docx':
        return None

    doc = Document(str(filepath))
    body_html = []

    # Iterate through document body elements in order (paragraphs and tables interleaved)
    for element in doc.element.body:
        tag_name = element.tag.split('}')[-1] if '}' in element.tag else element.tag

        if tag_name == 'p':
            # Find matching paragraph object
            for para in doc.paragraphs:
                if para._element is element:
                    text = para.text.strip()
                    if not text:
                        break

                    style_name = para.style.name if para.style else ''

                    if style_name.startswith('Heading'):
                        # Extract heading level
                        level_match = re.search(r'(\d+)', style_name)
                        level = int(level_match.group(1)) if level_match else 2
                        level = min(level, 6)
                        anchor = slugify(text)
                        inner = _para_to_html(para) or text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                        body_html.append(
                            f'<h{level} id="{anchor}">'
                            f'<a class="anchor-link" href="#{anchor}">#</a>'
                            f'{inner}</h{level}>'
                        )
                    else:
                        inner = _para_to_html(para)
                        if inner:
                            # Check for list styles
                            numPr = para._element.find(qn('w:pPr'))
                            is_list = numPr is not None and numPr.find(qn('w:numPr')) is not None
                            if is_list:
                                body_html.append(f'<li>{inner}</li>')
                            else:
                                body_html.append(f'<p>{inner}</p>')
                    break

        elif tag_name == 'tbl':
            for table in doc.tables:
                if table._tbl is element:
                    body_html.append(_table_to_html(table))
                    break

    title = filename.replace('.docx', '')

    return f'''<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{title} — Mommy's Heart</title>
    <style>
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        body {{
            font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
            background: #f8f9fa;
            color: #1a1a1a;
            line-height: 1.7;
        }}
        .content {{
            max-width: 800px;
            margin: 0 auto;
            padding: 32px 24px 80px;
            background: white;
            min-height: 100vh;
            box-shadow: 0 0 20px rgba(0,0,0,0.05);
        }}
        .doc-title {{
            font-size: 26px;
            font-weight: 700;
            color: #1a1a1a;
            margin-bottom: 24px;
            padding-bottom: 16px;
            border-bottom: 2px solid #1a73e8;
        }}
        h1, h2, h3, h4, h5, h6 {{
            margin: 28px 0 12px 0;
            color: #1a1a1a;
            position: relative;
        }}
        h1 {{ font-size: 24px; }}
        h2 {{ font-size: 20px; }}
        h3 {{ font-size: 17px; }}
        h4, h5, h6 {{ font-size: 15px; }}
        p {{ margin: 0 0 12px 0; font-size: 15px; }}
        li {{ margin: 0 0 6px 24px; font-size: 15px; }}
        strong {{ font-weight: 600; }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 16px 0;
            font-size: 14px;
        }}
        th, td {{
            border: 1px solid #dadce0;
            padding: 8px 12px;
            text-align: left;
        }}
        th {{ background: #f1f3f4; font-weight: 600; }}
        tr:nth-child(even) {{ background: #fafafa; }}
        .anchor-link {{
            color: #dadce0;
            text-decoration: none;
            font-weight: 400;
            margin-right: 6px;
            font-size: 0.8em;
        }}
        .anchor-link:hover {{ color: #1a73e8; }}
        /* Scroll target highlight */
        :target {{
            background: #fff3cd;
            padding: 4px 8px;
            border-radius: 4px;
            transition: background 2s;
        }}
    </style>
</head>
<body>
    <div class="content">
        <div class="doc-title">{title}</div>
        {"".join(body_html)}
    </div>
</body>
</html>'''
